"""物調表換算、Excel 產生、謄本解析、API 測試（不連網）。跑法：python -m unittest discover -s tests"""
import io
import json
import sys
import unittest
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from services import ai, geo, tenement  # noqa: E402
from sheet.excel import YELLOW, build_template, generate_excel  # noqa: E402
from sheet.layout import LAYOUT, cell_map, derive, missing_required  # noqa: E402
from sheet.schemas import PropertyForm  # noqa: E402

TODAY = date(2026, 1, 1)

# 虛構的謄本文字（格式仿電子謄本，內容全是假的）
FAKE_TENEMENT = """建物登記第二類謄本
示範區 示範段 00123-000建號
建物門牌 示範市示範路1號3樓
建物坐落地號 示範段
0123-0000
主要用途 集合住宅
主要建材 鋼筋混凝土造
層數 007層
層次 三層
總面積 100.00平方公尺
建築完成日期 民國095年06月15日
附屬建物用途 陽台 面積 10.00平方公尺
共有部分 示範段00200-000建號 200.00平方公 尺
權利範圍 100分之5
所有權人 測試甲
"""

SAMPLE = dict(case_number="A-001", address="示範市示範區示範路1號3樓", total_price=1280, case_type="華廈",
              build_area_main=30.25, build_area_sub=3.03, build_area_pub=3.03, room_count=3, hall_count=2,
              bath_count=2, finish_year=95, finish_month=6, finish_day=15, structure="鋼筋混凝土",
              floor_current=3, floor_above=7, visit_method="鑰匙", visit_key_no="12",
              land_section="示範段", land_number="0123-0000", land_zone_type="都計內", land_zone_name="住宅區",
              parking="有", parking_type="平面", management="管理員", mgmt_fee=1500, mgmt_period="月繳",
              orientation_build="南")


class LayoutTest(unittest.TestCase):
    def test_cell_map_unique_and_in_template(self):
        m = cell_map()
        self.assertEqual(len(set(m.values())), len(m))
        ws = build_template().active
        merged = {str(r).split(":")[0] for r in ws.merged_cells.ranges}
        for key, ref in m.items():
            self.assertIn(ref, merged, key)

    def test_derive_values(self):
        v = derive(PropertyForm(**SAMPLE), today=TODAY)
        self.assertEqual(v["build_area_total"], 36.31)
        self.assertEqual(v["unit_price"], 35.25)          # 1280 / 36.31
        self.assertEqual(v["layout"], "3房2廳2衛")
        self.assertEqual(v["floors"], "3 / 7 / 0")
        self.assertEqual(v["age_years"], 19.5)
        self.assertEqual(v["visit"], "鑰匙　鑰匙 #12")
        self.assertEqual(v["land_parcel"], "示範段 0123-0000 地號")
        self.assertEqual(v["land_use"], "都計內，住宅區")
        self.assertEqual(v["mgmt_text"], "管理員 1500 元／月繳")
        self.assertEqual(v["orientation"], "建物朝南")
        self.assertEqual(missing_required(v), [])

    def test_land_only_uses_land_area(self):
        v = derive(PropertyForm(case_number="L-1", address="示範段1地號", total_price=500, case_type="農地",
                                land_area_all=250), today=TODAY)
        self.assertEqual(v["unit_price"], 2)
        self.assertEqual(missing_required(v), [])

    def test_missing_required(self):
        v = derive(PropertyForm(address="示範路1號"), today=TODAY)
        self.assertIn("case_number", missing_required(v))
        self.assertIn("build_area_main", missing_required(v))


class ExcelTest(unittest.TestCase):
    def test_generate_and_read_back(self):
        form = PropertyForm(**SAMPLE)
        data = generate_excel(form, subtitle="示範不動產", today=TODAY)
        ws = load_workbook(io.BytesIO(data)).active
        m = cell_map()
        self.assertEqual(ws[m["case_number"]].value, "A-001")
        self.assertEqual(ws[m["total_price"]].value, 1280)
        self.assertEqual(ws[m["layout"]].value, "3房2廳2衛")
        self.assertEqual(ws["A2"].value, "示範不動產")
        self.assertEqual(form.missing_fields, [])

    def test_missing_marked_yellow(self):
        form = PropertyForm(address="示範路1號")
        ws = load_workbook(io.BytesIO(generate_excel(form, today=TODAY))).active
        ref = cell_map()["case_number"]
        self.assertEqual(ws[ref].fill.start_color.rgb[-6:], YELLOW.start_color.rgb[-6:])
        self.assertIn("case_number", form.missing_fields)

    def test_custom_template_file(self):
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "my.xlsx"
            wb = build_template("我的物調表")
            wb.save(p)
            ws = load_workbook(io.BytesIO(generate_excel(PropertyForm(**SAMPLE), p, today=TODAY))).active
            self.assertEqual(ws["A1"].value, "我的物調表")

    def test_layout_rows_match_sections(self):
        self.assertEqual(sum(1 for i in LAYOUT if i[0] == "section"), 5)


class TenementTest(unittest.TestCase):
    def test_parse_text(self):
        r = tenement.parse_text(FAKE_TENEMENT, today=TODAY)
        self.assertEqual(r["build_number"], "00123-000")
        self.assertEqual((r["land_section"], r["land_number"]), ("示範段", "0123-0000"))
        self.assertEqual(r["structure"], "鋼筋混凝土")
        self.assertEqual((r["floor_above"], r["floor_current"]), (7, 3))
        self.assertEqual(r["build_area_main"], 30.25)
        self.assertEqual(r["build_area_sub"], 3.02)
        self.assertEqual(r["build_area_pub"], 3.02)        # 200 × 5/100 = 10 m²
        self.assertEqual((r["finish_year"], r["finish_month"], r["finish_day"]), (95, 6, 15))
        self.assertNotIn("owner", r)                       # 不抓所有權人
        self.assertNotIn("測試甲", json.dumps(r, ensure_ascii=False))

    def test_zh_numbers(self):
        self.assertEqual([tenement.zh_to_int(x) for x in ("三", "十", "十二", "二十一", "12", "地下")],
                         [3, 10, 12, 21, 12, None])
        self.assertEqual(tenement.normalise_structure("鋼筋混凝土加強磚造"), "鋼筋混凝土加強磚造")


class GeoTest(unittest.TestCase):
    def test_parse_town_village(self):
        d = {"ctyName": "示範市", "townName": "示範區", "villageName": "示範里", "sectName": "示範段"}
        self.assertEqual(geo.parse_town_village(d),
                         {"county": "示範市", "town": "示範區", "village": "示範里", "land_section": "示範段"})
        self.assertEqual(geo.parse_town_village({"villName": "舊欄位里"})["village"], "舊欄位里")
        self.assertEqual(geo.parse_town_village([]), {})

    def test_tls_keeps_verification(self):
        import ssl
        ctx = geo._tls()
        self.assertEqual(ctx.verify_mode, ssl.CERT_REQUIRED)
        self.assertTrue(ctx.check_hostname)

    def test_nearest_places(self):
        els = [{"lat": 25.001, "lon": 121.0, "tags": {"leisure": "park", "name": "遠公園"}},
               {"lat": 25.0002, "lon": 121.0, "tags": {"leisure": "park", "name": "近公園"}},
               {"lat": 25.0, "lon": 121.001, "tags": {"railway": "station", "name": "示範站"}},
               {"lat": 25.0, "lon": 121.0, "tags": {"shop": "supermarket"}}]
        r = geo.nearest_places(els, 25.0, 121.0)
        self.assertTrue(r["park"].startswith("近公園"))
        self.assertTrue(r["station"].startswith("示範站"))
        self.assertEqual(r["market"], "")
        self.assertAlmostEqual(geo.distance_m(25, 121, 25.001, 121), 111.2, delta=0.5)


class AiFallbackTest(unittest.TestCase):
    def test_without_key(self):
        import os
        os.environ.pop("ANTHROPIC_API_KEY", None)
        self.assertEqual(ai.case_name({"address": "示範市中正區示範路1號", "room_count": 3, "case_type": "華廈"}),
                         "中正區3房華廈")
        self.assertIn("◆", ai.sell_points({"room_count": 2, "parking": "有"}))


class ApiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from fastapi.testclient import TestClient
        import main
        cls.c = TestClient(main.app)

    def test_pages_and_options(self):
        self.assertEqual(self.c.get("/").status_code, 200)
        self.assertEqual(self.c.get("/static/app.js").status_code, 200)
        self.assertIn("case_type", self.c.get("/api/dropdown-options").json())
        self.assertIn("brand", self.c.get("/api/settings").json())

    def test_generate_xlsx(self):
        r = self.c.post("/api/generate", data={"form_json": json.dumps({"address": "示範路1號"})})
        self.assertEqual(r.status_code, 200)
        self.assertIn("case_number", r.headers["x-missing-fields"])
        self.assertEqual(r.content[:2], b"PK")
        self.assertEqual(self.c.post("/api/generate", data={"form_json": "{\"total_price\": \"abc\"}"}).status_code, 400)

    def test_tenement_rejects_non_pdf(self):
        r = self.c.post("/api/parse-tenement", files={"file": ("x.pdf", b"hello", "application/pdf")})
        self.assertEqual(r.status_code, 400)


if __name__ == "__main__":
    unittest.main()
