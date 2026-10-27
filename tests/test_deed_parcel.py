"""謄本解析 v2、風險提示、地籍、門牌比對、照片頁測試（不連網）。內容全部是虛構的。"""
import asyncio
import base64
import io
import json
import sys
import unittest
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from services import parcel, risk, tenement  # noqa: E402
from sheet.excel import generate_excel  # noqa: E402
from sheet.layout import derive  # noqa: E402
from sheet.schemas import PropertyForm  # noqa: E402

TODAY = date(2026, 1, 1)

# 電傳版格式（標籤緊貼值）：虛構
LAND_DIANCHUAN = """土地登記第一類謄本（地號全部）
示範市示範區示範段0123-0000地號
土地標示部
登記日期民國080年01月02日登記原因地籍圖重測使用分區住宅區面積330.00平方公尺
民國114年01月公告土地現值50,000元/平方公尺
民國113年01月公告地價12,000元/平方公尺
土地所有權部
登記次序0001
登記日期民國110年03月04日
登記原因買賣
所有權人測試甲
統一編號A1*******
權利範圍100分之25
土地他項權利部
登記次序0001-000
權利種類最高限額抵押權
收件年期民國110年
權利人示範銀行股份有限公司
擔保債權總金額新臺幣3,600,000元正
查詢時間民國114年12月20日
"""

# 地政事務所正本格式（全形冒號、遮罩）：虛構、非都市土地、1 年內取得
LAND_OFFICIAL = """土地登記第二類謄本
示範縣示範鄉 示範段 0456-0001地號
土地標示部
登記日期：民國０９０年０５月０６日　　登記原因：分割
使用分區：一般農業區　　　　使用地類別：甲種建築用地
面　　積：＊＊＊＊1,000.00平方公尺
公告土地現值：民國114年01月　＊＊6,000元／平方公尺
土地所有權部
登記次序：０００２
登記日期：民國114年08月01日　　登記原因：贈與
所有權人：測試乙
權利範圍：全部　　　＊＊＊１分之１＊＊＊
列印時間：民國114年09月01日
"""

BUILDING_OFFICIAL = """建物登記第二類謄本
示範市示範區 示範段 01234-000建號
建物標示部
登記日期：民國１００年０１月０１日　登記原因：第一次登記
建物門牌：示範市示範區示範路１號５樓
建物坐落地號：示範段 0123-0000
主要用途：集合住宅
主要建材：鋼筋混凝土造
層　　數：０１２層
總面積：＊＊＊＊100.00平方公尺
層　　次：五層　　層次面積：＊＊＊＊100.00平方公尺
建築完成日期：民國０９９年１２月３１日
附屬建物用途：陽台　　面積：＊＊＊＊＊10.00平方公尺
共有部分：示範段01300-000建號 1,000.00平方公 尺
權利範圍：＊＊＊＊1000分之30＊＊＊＊
建物所有權部
登記次序：０００１
登記日期：民國１００年０２月０１日　登記原因：買賣
所有權人：測試丙
權利範圍：全部　＊＊＊１分之１
查詢時間：民國114年12月01日
"""


class LandDeedTest(unittest.TestCase):
    def test_dianchuan(self):
        r = tenement.parse_land_text(LAND_DIANCHUAN, TODAY)
        self.assertEqual((r["land_section"], r["land_number"]), ("示範段", "0123-0000"))
        self.assertEqual(r["land_zone_name"], "住宅區")
        self.assertEqual(r["land_zone_type"], "都計內")
        self.assertEqual(r["land_area_all"], 99.83)
        self.assertEqual(r["land_area_share"], 24.96)            # 99.825 × 25/100
        self.assertEqual((r["land_ownership"], r["land_share"]), ("持分", "100分之25"))
        self.assertEqual((r["land_value_per_sqm"], r["land_price_per_sqm"]), (50000, 12000))
        self.assertEqual((r["acquire_date"], r["acquire_reason"]), ("2021-03-04", "買賣"))
        self.assertNotIn("short_term", r)
        self.assertEqual(r["encumbrance_kinds"], ["最高限額抵押權"])
        self.assertEqual(r["encumbrance_amount"], 3600000)
        self.assertEqual(r["deed_date"], "2025-12-20")
        dump = json.dumps(r, ensure_ascii=False)
        for secret in ("測試甲", "A1", "示範銀行"):
            self.assertNotIn(secret, dump)                       # 不帶出姓名、統編、權利人

    def test_official_rural_short_term(self):
        r = tenement.parse_land_text(LAND_OFFICIAL, TODAY)
        self.assertEqual((r["land_section"], r["land_number"]), ("示範段", "0456-0001"))
        self.assertEqual((r["land_zone_type"], r["land_zone_name"], r["land_use_type"]),
                         ("都計外", "一般農業區", "甲種建築用地"))
        self.assertEqual(r["land_area_all"], 302.5)
        self.assertEqual((r["land_ownership"], r["land_area_share"]), ("全部", 302.5))
        self.assertEqual(r["land_value_per_sqm"], 6000)
        self.assertTrue(r["short_term"])
        self.assertEqual(r["acquire_reason"], "贈與")
        self.assertEqual(r["deed_date"], "2025-09-01")
        self.assertNotIn("encumbrance_kinds", r)
        self.assertNotIn("測試乙", json.dumps(r, ensure_ascii=False))


class BuildingDeedTest(unittest.TestCase):
    def test_official(self):
        r = tenement.parse_building_text(BUILDING_OFFICIAL, TODAY)
        self.assertEqual(r["build_number"], "01234-000")
        self.assertEqual((r["land_section"], r["land_number"]), ("示範段", "0123-0000"))
        self.assertEqual((r["main_use"], r["structure"]), ("集合住宅", "鋼筋混凝土"))
        self.assertEqual((r["floor_above"], r["floor_current"]), (12, 5))
        self.assertEqual(r["layers"], [{"floor": "五層", "ping": 30.25}])
        self.assertEqual((r["build_area_main"], r["build_area_sub"], r["build_area_pub"]), (30.25, 3.02, 9.07))
        self.assertEqual(r["sub_uses"], ["陽台"])
        self.assertEqual((r["finish_year"], r["finish_month"], r["finish_day"]), (99, 12, 31))
        self.assertEqual(r["build_share"], "全部")
        self.assertEqual(r["acquire_date"], "2011-02-01")
        self.assertEqual(r["deed_date"], "2025-12-01")
        self.assertNotIn("測試丙", json.dumps(r, ensure_ascii=False))

    def test_not_a_deed(self):
        self.assertEqual(tenement.parse_building_text("隨便一段文字"), {})
        self.assertEqual(tenement.parse_land_text("隨便一段文字"), {})

    def test_parse_pdf_detects_kind(self):
        try:
            import fitz
        except ImportError:
            self.skipTest("沒裝 PyMuPDF")
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((40, 60), "土地標示部\n示範段0123-0000地號\n面積100.00平方公尺", fontname="china-t", fontsize=11)
        r = tenement.parse_pdf(doc.tobytes(), TODAY)
        self.assertEqual(r["kind"], "land")
        self.assertEqual(r["land_area_all"], 30.25)


class RiskTest(unittest.TestCase):
    def test_red_flags(self):
        deed = tenement.parse_land_text(LAND_OFFICIAL, TODAY)
        deed.update({"encumbrance_kinds": ["抵押權"], "encumbrance_amount": 1000000})
        items = risk.check(deed, {}, TODAY)
        levels = {i["item"]: i["level"] for i in items}
        self.assertEqual(levels["短期交易"], "red")
        self.assertEqual(levels["他項權利"], "red")
        self.assertEqual(levels["謄本鮮度"], "yellow")           # 2025-09-01 → 122 天
        self.assertTrue(risk.summary(items).startswith("🔴 2"))
        self.assertIn("1,000,000", risk.text_for_sheet(items))

    def test_area_mismatch_and_share(self):
        deed = {"acquire_date": "2015-01-01", "holding_years": 11, "deed_date": "2025-12-25",
                "build_area_total": 40.0, "land_area_share": 20.0, "land_ownership": "持分"}
        items = risk.check(deed, {"claimed_build_area_total": 42.0, "claimed_land_area_share": 20.2}, TODAY)
        levels = {i["item"]: i["level"] for i in items}
        self.assertEqual(levels["坪數核對（建物權狀坪數）"], "red")   # 5%
        self.assertEqual(levels["坪數核對（土地持分坪數）"], "green")  # 1%
        self.assertEqual(levels["持分"], "yellow")
        self.assertEqual(levels["他項權利"], "green")

    def test_no_deed(self):
        self.assertEqual(risk.check({}, {}, TODAY)[0]["level"], "yellow")


def fake_mask(w=10, h=10, fill=50):
    """產生一張 w×h 的 PNG，前 fill% 的像素不透明（alpha 127），模擬圖台遮罩。"""
    from PIL import Image
    im = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    n = int(w * h * fill / 100)
    for i in range(n):
        im.putpixel((i % w, i // w), (255, 0, 0, 127))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return base64.b64encode(buf.getvalue()).decode()


class ParcelTest(unittest.TestCase):
    def test_landno8_and_section(self):
        self.assertEqual([parcel.landno8(x) for x in ("401", "401-1", "401之1", "0401-0001", "４０１")],
                         ["04010000", "04010001", "04010001", "04010001", "04010000"])
        secs = [("XA", "0001", "示範段"), ("XA", "0002", "示範段一小段"), ("XA", "0003", "東示範段")]
        self.assertEqual(parcel.pick_section(secs, "示範段")[1], "0001")
        self.assertEqual(parcel.pick_section(secs, "示範段一小段")[1], "0002")
        self.assertIsNone(parcel.pick_section(secs, "不存在段"))

    def test_parse_rows_area_and_dims(self):
        try:
            img = fake_mask(fill=50)
        except ImportError:
            self.skipTest("沒裝 Pillow")
        # 外框約 100m × 100m，填一半 → 約 5,000 m²
        rows = [{"sect": "0001", "landno": "00010000"},
                {"image": img, "lx": "121.0", "ly": "24.0", "rx": str(121 + 100 / (111320 * 0.9135)), "ry": str(24 + 100 / 110574)}]
        r = parcel.parse_parcel_rows(rows)
        self.assertAlmostEqual(r["area_sqm"], 5000, delta=60)
        self.assertAlmostEqual(r["width_m"], 100, delta=1)
        self.assertAlmostEqual(r["depth_m"], 100, delta=1)
        self.assertEqual(parcel.parse_parcel_rows([{"msg": "查無資料"}]), {})

    def test_multi_part_parcel_sums_area(self):
        try:
            a, b = fake_mask(fill=100), fake_mask(fill=20)
        except ImportError:
            self.skipTest("沒裝 Pillow")
        d = 50 / 110574
        rows = [{"image": a, "lx": 121.0, "ly": 24.0, "rx": 121.0 + d, "ry": 24.0 + d},
                {"image": b, "lx": 121.001, "ly": 24.0, "rx": 121.001 + d, "ry": 24.0 + d}]
        r = parcel.parse_parcel_rows(rows)
        self.assertEqual(len(r["parts"]), 2)
        self.assertAlmostEqual(r["area_sqm"], sum(p["area_sqm"] for p in r["parts"]), delta=0.2)

    def test_address_normalise_and_core(self):
        self.assertEqual(parcel.normalise_address("示範市示範區示範里３鄰示範路１２號５樓"), "示範市示範區示範路12號")
        self.assertEqual(parcel.normalise_address("示範鎮示範新村12號"), "示範鎮示範新村12號")  # 社區名保留
        self.assertEqual(parcel.address_core("示範市示範區示範里4鄰示範路一段12巷3號5樓"), "示範路一段12巷3號")

    def test_query_search_strict(self):
        xml = """<root><ITEM><CONTENT>示範市示範區示範里４鄰示範路１２號</CONTENT><KEY>A,ADDRESS,1</KEY>
<LOCATION>121.5,25.0</LOCATION></ITEM><ITEM><CONTENT>示範市示範區示範路１４號</CONTENT>
<KEY>A,ADDRESS,2</KEY><LOCATION>121.6,25.1</LOCATION></ITEM></root>"""
        hit = parcel.parse_query_search(xml, parcel.address_core("示範市示範區示範路12號"))
        self.assertEqual((hit["lat"], hit["lon"]), (25.0, 121.5))
        self.assertIsNone(parcel.parse_query_search(xml, parcel.address_core("示範市示範區示範路99號")))

    def test_resolve_section_with_fake_lists(self):
        async def fake_counties():
            return [("X", "示範市")]

        async def fake_towns(cc):
            return [("X01", "東區"), ("X02", "西區")]

        async def fake_sections(cc, tc):
            return [("XA", "0101", "示範段")] if tc == "X02" else [("XA", "0201", "別的段")]
        saved = parcel.counties, parcel.towns, parcel.sections
        parcel.counties, parcel.towns, parcel.sections = fake_counties, fake_towns, fake_sections
        try:
            self.assertEqual(asyncio.run(parcel.resolve_section("示範市", "示範段")), ("XA", "0101", "西區", "示範段"))
            with self.assertRaises(ValueError):
                asyncio.run(parcel.resolve_section("不存在市", "示範段"))
        finally:
            parcel.counties, parcel.towns, parcel.sections = saved


class SheetV2Test(unittest.TestCase):
    def test_derive_rights_and_risk(self):
        f = PropertyForm(land_ownership="持分", land_share="100分之25", land_value_per_sqm=50000,
                         acquire_date="2021-03-04", acquire_reason="買賣", holding_years=4.8,
                         encumbrance_kinds=["最高限額抵押權"], encumbrance_amount=3600000, risk_notes="🔴 測試")
        v = derive(f, today=TODAY)
        self.assertEqual(v["land_share_text"], "持分 100分之25")
        self.assertEqual(v["land_value_text"], "50,000 元/㎡（約 16.53 萬/坪）")
        self.assertEqual(v["acquire_text"], "2021-03-04　買賣　持有約 4.8 年")
        self.assertEqual(v["encumbrance_text"], "最高限額抵押權（擔保約 360 萬）")
        self.assertEqual(v["risk_notes"], "🔴 測試")

    def test_photos_sheet_strips_exif(self):
        try:
            from PIL import Image
        except ImportError:
            self.skipTest("沒裝 Pillow")
        from openpyxl import load_workbook
        im = Image.new("RGB", (1600, 1200), (200, 180, 160))
        exif = Image.Exif()
        exif[0x010F] = "DemoCamera"       # Make
        buf = io.BytesIO()
        im.save(buf, "JPEG", exif=exif)
        data = generate_excel(PropertyForm(address="示範路1號"), today=TODAY,
                              photos=[("外觀.jpg", buf.getvalue()), ("客廳.jpg", buf.getvalue())])
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["物調表", "照片"])
        self.assertEqual(len(wb["照片"]._images), 2)
        self.assertEqual(wb["照片"]["A18"].value, "外觀.jpg")
        self.assertNotIn(b"DemoCamera", data)
        import zipfile
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for n in z.namelist():
                if n.startswith("xl/media/"):
                    self.assertNotIn(b"Exif", z.read(n)[:200])


class ApiV2Test(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from fastapi.testclient import TestClient
        import main
        cls.c = TestClient(main.app)

    def test_risk_endpoint(self):
        deed = tenement.parse_land_text(LAND_DIANCHUAN, TODAY)
        r = self.c.post("/api/risk", data={"deed_json": json.dumps(deed), "form_json": "{}"}).json()
        self.assertIn("summary", r)
        self.assertTrue(any(i["item"] == "他項權利" and i["level"] == "red" for i in r["items"]))

    def test_generate_with_bad_photo(self):
        r = self.c.post("/api/generate", data={"form_json": "{}"},
                        files=[("photos", ("x.jpg", b"not an image", "image/jpeg"))])
        self.assertEqual(r.status_code, 400)


if __name__ == "__main__":
    unittest.main()
