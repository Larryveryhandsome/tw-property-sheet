"""產生通用物調表範本，並把資料填進去。

範本是用程式畫出來的（不依賴任何品牌的既有表單），A4 直式一頁。
"""
import io
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .layout import LAYOUT, cell_map, derive, missing_required

YELLOW = PatternFill("solid", start_color="FFF59D", end_color="FFF59D")
LABEL_FILL = PatternFill("solid", start_color="EEF3F6", end_color="EEF3F6")
SECTION_FILL = PatternFill("solid", start_color="2F5D73", end_color="2F5D73")
THIN = Side(style="thin", color="9AA9B2")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="center", wrap_text=True)


def safe_set(ws, ref, v):
    """寫入儲存格。開頭是 = + - @ 的文字一律當純文字（防 Excel 公式注入：資料可能來自表單、謄本或 OSM 地名）。"""
    ws[ref] = v
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        ws[ref].data_type = "s"
        ws[ref].quotePrefix = True


def build_template(title: str = "物件調查表", subtitle: str = "") -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "物調表"
    for col, width in zip("ABCDEF", (16, 15, 15, 16, 15, 15)):
        ws.column_dimensions[col].width = width
    ws.merge_cells("A1:F1")
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:F2")
    ws["A2"] = subtitle
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A2"].font = Font(size=9, color="666666")

    row = 3
    for item in LAYOUT:
        kind = item[0]
        if kind == "section":
            ws.merge_cells(f"A{row}:F{row}")
            c = ws[f"A{row}"]
            c.value, c.fill = item[1], SECTION_FILL
            c.font = Font(bold=True, color="FFFFFF")
        elif kind == "row":
            for (key, label), (lc, v1, v2) in zip(item[1], (("A", "B", "C"), ("D", "E", "F"))):
                ws[f"{lc}{row}"] = label
                ws.merge_cells(f"{v1}{row}:{v2}{row}")
        else:  # wide / text
            ws[f"A{row}"] = item[2]
            ws.merge_cells(f"B{row}:F{row}")
            if kind == "text":
                ws.row_dimensions[row].height = item[3]
        if kind != "section":
            for col in "ABCDEF":
                cell = ws[f"{col}{row}"]
                cell.border = BOX
                cell.alignment = WRAP
            for col in ("A", "D") if kind == "row" else ("A",):
                ws[f"{col}{row}"].fill = LABEL_FILL
                ws[f"{col}{row}"].font = Font(size=10, bold=True)
        if kind != "text":
            ws.row_dimensions[row].height = ws.row_dimensions[row].height or 20
        row += 1

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:F{row - 1}"
    ws.oddFooter.center.text = "本表資料原則上依現況、謄本與公開資料整理，實際以登記資料與現場確認為準"
    return wb


def add_photos(wb, photos, max_px=900):
    """照片放第二張工作表：每張縮到長邊 max_px，一列兩張，下方寫檔名。需要 Pillow。"""
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image, ImageOps
    ws = wb.create_sheet("照片")
    ws.column_dimensions["A"].width = ws.column_dimensions["B"].width = 62
    row = 1
    for i, (name, data) in enumerate(photos):
        im = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("RGB")
        im.thumbnail((max_px, max_px))
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=85)   # 重新存檔＝不帶原圖 EXIF（GPS、拍攝裝置）
        buf.seek(0)
        xl = XLImage(buf)
        scale = 430 / max(xl.width, 1)      # 欄寬約 430 像素
        xl.width, xl.height = int(xl.width * scale), int(xl.height * scale)
        col = "A" if i % 2 == 0 else "B"
        ws.add_image(xl, f"{col}{row}")
        safe_set(ws, f"{col}{row + 17}", name)
        if i % 2 == 1:
            row += 19
    for r in range(1, row + 19):
        ws.row_dimensions[r].height = 18
    return ws


def generate_excel(form, template_path: Path = None, title="物件調查表", subtitle="", today=None, photos=None) -> bytes:
    """回傳 xlsx bytes。有指定 template_path 就用它（版面要和 LAYOUT 一致）。photos：[(檔名, bytes)]。"""
    if template_path and Path(template_path).exists():
        wb = load_workbook(template_path)
    else:
        wb = build_template(title, subtitle)
    ws = wb.worksheets[0]
    values = derive(form, today=today)
    cmap = cell_map()
    for key, ref in cmap.items():
        v = values.get(key)
        if v not in (None, ""):
            safe_set(ws, ref, v)
    missing = missing_required(values)
    for key in missing:
        ws[cmap[key]].fill = YELLOW
    if hasattr(form, "missing_fields"):
        form.missing_fields = missing
    if photos:
        add_photos(wb, photos)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
